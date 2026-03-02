"""
generate_sample_data.py
-----------------------
Generates a sample telemetry Excel workbook (data/sample_telemetry.xlsx)
with the column layout expected by the Power BI dashboard.

Column mapping (1-based, A=1):
  Col 1  (A) – LAP           : lap number
  Col 2  (B) – (reserved)
  Col 3  (C) – (reserved)
  Col 4  (D) – LAPTIME       : raw MM.SS.mmm  (e.g. 1.23.456 = 1 min 23.456 s)
  Col 5  (E) – FUEL/LAP      : fuel consumed per lap (L)
  Col 6  (F) – T wat, Max    : max water temperature (°C)
  Col 7  (G) – T wat, avg    : avg water temperature (°C)
  Col 8  (H) – T oil, Max    : max oil temperature (°C)
  Col 9  (I) – T oil, avg    : avg oil temperature (°C)
  Col 10 (J) – Poil, Min     : min oil pressure (bar)
  Col 11 (K) – Poil, avg     : avg oil pressure (bar)
  Col 12 (L) – Alt V, Min    : min alternator voltage (V)
  Col 13 (M) – Alt V, avg    : avg alternator voltage (V)
  Col 14 (N) – Pfuel, Min    : min fuel pressure (bar)
  Col 15 (O) – Pfuel, avg    : avg fuel pressure (bar)
  Col 16 (P) – LiftPump1 current, avg : avg lift-pump 1 current (A)
  Col 17 (Q) – LiftPump2 current, avg : avg lift-pump 2 current (A)

Row 1 : main header row
Row 2 : unit / sub-header row
Row 3+: data rows
"""

import random
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_FILE = os.path.join(REPO_ROOT, "data", "sample_telemetry.xlsx")

HEADERS = [
    "LAP", "", "", "LAPTIME",
    "FUEL/LAP",
    "T wat, Max", "T wat, avg",
    "T oil, Max", "T oil, avg",
    "Poil, Min", "Poil, avg",
    "Alt V, Min", "Alt V, avg",
    "Pfuel, Min", "Pfuel, avg",
    "LiftPump1 current, avg", "LiftPump2 current, avg",
]

UNITS = [
    "#", "", "", "MM.SS.mmm",
    "L",
    "°C", "°C",
    "°C", "°C",
    "bar", "bar",
    "V", "V",
    "bar", "bar",
    "A", "A",
]

NUM_LAPS = 20
random.seed(42)


def fmt_laptime(total_seconds: float) -> str:
    """Convert total seconds (SSS.mmm) to MM.SS.mmm string."""
    minutes = int(total_seconds // 60)
    seconds = total_seconds - minutes * 60
    sec_int = int(seconds)
    millis = round((seconds - sec_int) * 1000)
    return f"{minutes}.{sec_int:02d}.{millis:03d}"


def generate_laps(n: int) -> list[dict]:
    laps = []
    base_lap = 83.5   # ~1:23.500
    for i in range(1, n + 1):
        lap_time = base_lap + random.gauss(0, 0.8)
        laps.append({
            "lap": i,
            "laptime_raw": fmt_laptime(max(lap_time, 75.0)),
            "fuel_per_lap": round(random.uniform(2.8, 3.4), 3),
            "t_wat_max": round(random.uniform(88, 96), 1),
            "t_wat_avg": round(random.uniform(82, 90), 1),
            "t_oil_max": round(random.uniform(105, 118), 1),
            "t_oil_avg": round(random.uniform(98, 112), 1),
            "poil_min": round(random.uniform(3.2, 4.0), 2),
            "poil_avg": round(random.uniform(4.0, 5.0), 2),
            "altv_min": round(random.uniform(13.5, 14.0), 2),
            "altv_avg": round(random.uniform(14.0, 14.5), 2),
            "pfuel_min": round(random.uniform(3.8, 4.2), 2),
            "pfuel_avg": round(random.uniform(4.2, 4.8), 2),
            "lp1_avg": round(random.uniform(5.0, 6.5), 2),
            "lp2_avg": round(random.uniform(5.0, 6.5), 2),
        })
    return laps


def write_workbook(laps: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Telemetry"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    unit_fill = PatternFill("solid", fgColor="2E75B6")
    header_font = Font(bold=True, color="FFFFFF")

    # Row 1 – headers
    for col_idx, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Row 2 – units
    for col_idx, unit in enumerate(UNITS, start=1):
        cell = ws.cell(row=2, column=col_idx, value=unit)
        cell.font = Font(italic=True, color="FFFFFF")
        cell.fill = unit_fill
        cell.alignment = Alignment(horizontal="center")

    # Rows 3+ – data
    for row_offset, lap in enumerate(laps):
        r = 3 + row_offset
        ws.cell(row=r, column=1, value=lap["lap"])
        # cols 2-3 intentionally empty
        ws.cell(row=r, column=4, value=lap["laptime_raw"])
        ws.cell(row=r, column=5, value=lap["fuel_per_lap"])
        ws.cell(row=r, column=6, value=lap["t_wat_max"])
        ws.cell(row=r, column=7, value=lap["t_wat_avg"])
        ws.cell(row=r, column=8, value=lap["t_oil_max"])
        ws.cell(row=r, column=9, value=lap["t_oil_avg"])
        ws.cell(row=r, column=10, value=lap["poil_min"])
        ws.cell(row=r, column=11, value=lap["poil_avg"])
        ws.cell(row=r, column=12, value=lap["altv_min"])
        ws.cell(row=r, column=13, value=lap["altv_avg"])
        ws.cell(row=r, column=14, value=lap["pfuel_min"])
        ws.cell(row=r, column=15, value=lap["pfuel_avg"])
        ws.cell(row=r, column=16, value=lap["lp1_avg"])
        ws.cell(row=r, column=17, value=lap["lp2_avg"])

    # Auto-fit column widths (approximate)
    for col_idx, header in enumerate(HEADERS, start=1):
        max_len = max(len(header), 8)
        ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 4

    os.makedirs(os.path.dirname(OUT_FILE), exist_ok=True)
    wb.save(OUT_FILE)
    print(f"Saved: {OUT_FILE}")


if __name__ == "__main__":
    laps = generate_laps(NUM_LAPS)
    write_workbook(laps)
